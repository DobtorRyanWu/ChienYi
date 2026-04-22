# -*- coding: utf-8 -*-
"""A 標自主檢查「總表單」xlsx 解析器

輸入：`自檢表總表單0308.xlsx` 之類的多 sheet 彙總檔
輸出：list[dict]，每個 dict 對應一列檢查事件，可餵給
      general.self.inspection.sudo().create()

總表單結構（29 個 category sheet 共通）：
    R1  臺北市政府工務局水利工程處 ... (C6: 檔案路徑)
    R2  磺港溪再造C段... (C7: sheet label)
    R3  XXX 抽查表統計
    R4  header: 次數 | 編號 | 檢查日期 | 檢查部位 | 備註 | (C6, C7 是輔助欄)
    R5+ data rows

Sheet 命名規則：`^(\\d+)\\.(.+)$`，例 `1.測量放樣`、`6.混凝土`、`29.既有橋梁拆除`。
非此格式的 sheet（統計圖、鋼筋取樣、PC樁、混凝土取樣、混凝土送驗單）跳過。
"""

import re
from datetime import date, datetime
from io import BytesIO

import openpyxl


_SHEET_NAME_RE = re.compile(r'^(\d+)\.(.+)$')


def _cell(ws, r, c):
    try:
        return ws.cell(row=r, column=c).value
    except Exception:
        return None


def _strip(v):
    if v is None:
        return ''
    return str(v).strip()


def _coerce_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def parse_inspection_summary_xlsm(file_bytes: bytes, filename: str = '') -> list:
    """解析自檢表總表單，回傳 list[dict]。

    每個 dict：
        {
            'sheet_label': '1.測量放樣',
            'type_name': '測量放樣',
            'seq_no': 1,
            'inspection_date': date(2023, 11, 13),
            'location': '新設河道(0K+300~0K+450)',
            'note': '',
        }

    raise ValueError 當檔案無法開啟。
    """
    if not file_bytes:
        raise ValueError(f'檔案 {filename} 內容為空')

    try:
        wb = openpyxl.load_workbook(BytesIO(file_bytes), data_only=True, read_only=False)
    except Exception as e:
        raise ValueError(f'檔案 {filename} 無法開啟：{e}')

    rows = []
    for sheet_name in wb.sheetnames:
        m = _SHEET_NAME_RE.match(sheet_name)
        if not m:
            continue
        seq_prefix = int(m.group(1))
        type_name = m.group(2).strip()
        ws = wb[sheet_name]

        for r in range(5, ws.max_row + 1):
            raw_date = _cell(ws, r, 3)
            d = _coerce_date(raw_date)
            if not d:
                continue
            location = _strip(_cell(ws, r, 4))
            note = _strip(_cell(ws, r, 5))
            seq_no = _cell(ws, r, 1)
            try:
                seq_no = int(seq_no) if seq_no is not None else 0
            except (TypeError, ValueError):
                seq_no = 0

            rows.append({
                'sheet_label': sheet_name,
                'type_seq': seq_prefix,
                'type_name': type_name,
                'seq_no': seq_no,
                'inspection_date': d,
                'location': location,
                'note': note,
            })

    return rows
