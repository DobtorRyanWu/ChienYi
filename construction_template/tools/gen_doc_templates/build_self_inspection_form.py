# -*- coding: utf-8 -*-
"""產生「自主檢查表（單張）」的系統預設空白樣板。

輸入：使用者提供的 inspection_template.docx（已 Jinja2 化的草稿）
輸出：data/templates_blank/self_inspection_form.docx

草稿有三個必須修的問題（2026-09-11 實測）：

1. 表格裡的 `{% for %}` / `{% endfor %}` 寫成普通 Jinja2 標籤，而它們各自
   佔一整列。docxtpl 只有 `{%tr %}` 才會把整列刪掉，所以產出會在每個項目
   之間夾一個空白列（3 段 9 項 → 34 列，其中 16 列空白）。改成 `{%tr %}`
   之後是 18 列，9 項正好 9 列。

2. 照片頁的外層 `{% for image in inspection.images %}` 包住整頁，表格裡又有
   一個同名的 `{% for image %}`——巢狀同名迴圈會產生 N×N 頁；而且
   `{{ image.__self__ }}` 不是 docxtpl 輸出圖片的方式（要 InlineImage）。
   改成一頁兩張、由對照表先切好 imagePages。

3. 沒有量測區塊（表尾「丈量___位置…□合格□不合格」那一段）。紙本上它在
   檢查項目表格的最後、缺失複查之前。

用法（專案根目錄）：
    python construction_template/tools/gen_doc_templates/build_self_inspection_form.py \
           <來源 inspection_template.docx>
"""

import copy
import os
import re
import sys

import docx as python_docx
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, '..', '..', 'data', 'templates_blank',
                   'self_inspection_form.docx')

LOOP_TAG = re.compile(r'^\{%\s*(for\b.*?|endfor)\s*%\}$')


# ---------------------------------------------------------------- 小工具
def cell_text(cell):
    return '\n'.join(p.text for p in cell.paragraphs).strip()


def set_cell_text(cell, text):
    """把儲存格內容換成單一段落單一 run（保留第一段的樣式）"""
    para = cell.paragraphs[0]
    for run in list(para.runs)[1:]:
        run._element.getparent().remove(run._element)
    if para.runs:
        para.runs[0].text = text
    else:
        para.add_run(text)
    for extra in list(cell.paragraphs)[1:]:
        extra._element.getparent().remove(extra._element)


def is_full_span_row(row):
    """整列合併成一格（cells 都指向同一個 tc）"""
    return len({id(c._tc) for c in row.cells}) == 1


def row_text(row):
    return cell_text(row.cells[0]) if row.cells else ''


# ---------------------------------------------------------------- 步驟 1
def fix_tr_tags(doc):
    """表格列上的 {% for %} / {% endfor %} → {%tr %}"""
    n = 0
    for table in doc.tables:
        for row in table.rows:
            if not is_full_span_row(row):
                continue
            m = LOOP_TAG.match(row_text(row))
            if not m:
                continue
            set_cell_text(row.cells[0], '{%%tr %s %%}' % m.group(1).strip())
            n += 1
    return n


# ---------------------------------------------------------------- 步驟 2
MEASURE_ROWS = [
    '{%tr for mb in inspection.measures %}',
    '{{ mb.title }}',
    '{%tr for m in mb.lines %}',
    '{{ m.no }}.{{ m.text }}　{{ m.passMark }}合格　{{ m.failMark }}不合格',
    '{%tr endfor %}',
    '{%tr endfor %}',
]


def insert_measure_block(doc):
    """在「缺失複查結果」那一列之前插入量測區塊的 6 列。

    以該列（整列合併、跨全欄）為版型範本複製，才會沿用樣板自己的框線與字型。
    """
    table = doc.tables[0]
    anchor = None
    for row in table.rows:
        if is_full_span_row(row) and row_text(row).startswith('缺失複查結果'):
            anchor = row
            break
    if anchor is None:
        raise SystemExit('找不到「缺失複查結果」那一列，無法決定量測區塊的插入位置')

    for text in MEASURE_ROWS:
        new_tr = copy.deepcopy(anchor._tr)
        anchor._tr.addprevious(new_tr)
        from docx.table import _Row
        set_cell_text(_Row(new_tr, table).cells[0], text)
    return len(MEASURE_ROWS)


# ---------------------------------------------------------------- 步驟 3
def rebuild_photo_page(doc):
    """重做照片頁：一頁兩張，由對照表先切好 inspection.imagePages。"""
    body = doc.element.body
    children = list(body)

    # 找出照片頁的範圍：從含 `{% for image` 的段落，到含 `{% endfor %}` 的段落
    start = end = None
    for i, el in enumerate(children):
        if el.tag != qn('w:p'):
            continue
        text = ''.join(t.text or '' for t in el.iter(qn('w:t')))
        if start is None and '{% for image' in text:
            start = i
        elif start is not None and '{% endfor %}' in text:
            end = i
    if start is None or end is None:
        raise SystemExit('找不到照片頁的迴圈範圍')

    # 外層迴圈改成走 imagePages；{%p %} 讓標記所在的段落自己消失
    def set_para(el, text):
        for run in list(el.findall(qn('w:r')))[1:]:
            el.remove(run)
        runs = el.findall(qn('w:r'))
        if runs:
            for t in runs[0].iter(qn('w:t')):
                t.text = text
                break
            else:
                pass
        return el

    set_para(children[start], '{%p for page in inspection.imagePages %}')
    set_para(children[end], '{%p endfor %}')

    # 照片表格：草稿是 3 列（迴圈標記 + 內容 + 迴圈標記），改成 2 列各放一張
    photo_table = doc.tables[1]
    rows = photo_table.rows
    if len(rows) != 3:
        raise SystemExit('照片表格預期 3 列，實際 %d 列' % len(rows))

    tmpl_tr = copy.deepcopy(rows[1]._tr)          # 中間那列是內容列
    for row in list(rows):
        row._tr.getparent().remove(row._tr)

    from docx.table import _Row
    for slot in ('first', 'second'):
        new_tr = copy.deepcopy(tmpl_tr)
        photo_table._tbl.append(new_tr)
        cell = _Row(new_tr, photo_table).cells[0]
        if slot == 'first':
            body_text = ('{{ page.first.image }}\n'
                         '拍攝日期：{{ page.first.date }}　'
                         '說明：{{ page.first.description }}')
        else:
            body_text = ('{% if page.second %}{{ page.second.image }}\n'
                         '拍攝日期：{{ page.second.date }}　'
                         '說明：{{ page.second.description }}{% endif %}')
        set_cell_text(cell, body_text)
    return 2


def main():
    src = sys.argv[1]
    doc = python_docx.Document(src)
    n_tr = fix_tr_tags(doc)
    n_measure = insert_measure_block(doc)
    n_photo = rebuild_photo_page(doc)
    out = os.path.normpath(OUT)
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)
    print('迴圈標記改 {%%tr %%}：%d 處' % n_tr)
    print('量測區塊插入：%d 列' % n_measure)
    print('照片頁重建：%d 列' % n_photo)
    print('輸出：%s' % out)


if __name__ == '__main__':
    main()
