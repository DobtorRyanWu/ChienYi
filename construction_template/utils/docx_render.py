# -*- coding: utf-8 -*-
"""docx 樣板套印。

舊 EAGLE 樣板的 docx 用 `+++...+++` 語法：

    +++INS name+++            取值
    +++INS $item.field+++     迴圈變數取值
    +++FOR v IN col+++        段落層級迴圈
    +++FOR v IN $col+++       表格列層級迴圈
    +++END-FOR v+++           迴圈結束
    +++IF x+++ … +++END-IF+++ 條件（只有檢試驗管制用到）

`dobtor_doc_editor` 已經實作了 INS / FOR / END-FOR → docxtpl Jinja2 的轉換
（`_convert_ins_to_jinja`，段落層級操作，正確處理 Word 把標記切碎成多個
`<w:r>` run 的問題），直接沿用不重寫。

**但它不處理 `+++IF+++` / `+++END-IF+++`**（實測 2026-08-05：`_ANY` 只認
INS|FOR|END-FOR），未處理的標記會原樣印在報表上。本模組補上這一段——
採「移除標記、保留內容」：對照表保證條件欄位一定存在（取不到就給空字串），
所以條件本身沒有意義，留著反而會印出垃圾字。
"""

import io
import logging
import re

_logger = logging.getLogger(__name__)

# +++IF <任意運算式>+++ / +++END-IF+++
# 條件式是舊系統的 JS 運算式（例：$record.images[0].isEmpty !== true），
# 含 [] ! = 等字元，不能只認識別字元。
IF_RE = re.compile(r'\+{3}IF\s+[^+]{1,120}?\+{3,4}')
ENDIF_RE = re.compile(r'\+{3}END-IF\s*\+{3,4}')

# +++IMAGE imageGenerator($record.images[0].src, $record.images[0].extension,
#                         {height:6}, $record.images[0].date)+++
# 抓出「哪個迴圈變數的第幾張圖」，轉成我們自己的 context 鍵。
IMAGE_RE = re.compile(
    r'\+{3}IMAGE\s+imageGenerator\(\s*\$?(\w+)\.images\[(\d+)\]\.src[^+]*?\+{3,4}')
IMAGE_KEY = '__image__'          # context 裡的圖片佔位標記

# 帶索引的取值：+++INS $record.images[0].description+++
# 既有轉換器的 _INS 只認 `[A-Za-z_][\w]*` 加點號，不支援 [0]，
# 這類標記會原樣留在報表上，所以自己先轉掉。
INDEXED_INS_RE = re.compile(
    r'\+{3}INS\s+\$?\s*'
    r'([A-Za-z_]\w*(?:\s*\.\s*[A-Za-z_]\w*|\s*\[\s*\d+\s*\])*'
    r'\s*\[\s*\d+\s*\]'
    r'(?:\s*\.\s*[A-Za-z_]\w*|\s*\[\s*\d+\s*\])*)\s*\+{3,4}')


def _strip_conditionals(raw_bytes):
    """移除 +++IF+++ / +++END-IF+++ 標記，保留中間內容。

    與既有轉換器同樣在段落層級操作（para.text 已拼接所有 run），
    才不會被 Word 的 run 切割弄壞。
    """
    import docx as python_docx

    doc = python_docx.Document(io.BytesIO(raw_bytes))

    def clean(paragraphs):
        for para in paragraphs:
            text = para.text
            if '+++' not in text:
                continue
            # 相鄰標記會黏成一長串加號：`+++END-IF++++++INS x+++` 中間是 6 個
            # （前一個收尾 3~4 個 ＋ 下一個開頭 3 個）。既有轉換器的正則結尾吃
            # `\+{3,4}`，會啃掉下一個標記的開頭，導致 +++INS 原樣留在報表上。
            # 用零寬空格切開，輸出看不見但正則能正確斷詞。
            text = re.sub(r'\+{6,}', '+++​+++', text)
            # 圖片：轉成自己的 context 鍵，由 render() 換成 docxtpl 的 InlineImage
            text = IMAGE_RE.sub(
                lambda m: '{{ %s.images[%s].image }}' % (m.group(1), m.group(2)), text)
            # 帶索引的取值：既有轉換器處理不了，自己轉成 Jinja
            text = INDEXED_INS_RE.sub(
                lambda m: '{{ %s }}' % re.sub(r'\s', '', m.group(1)), text)
            text = ENDIF_RE.sub('', IF_RE.sub('', text))
            # 整段重寫成單一 run：標記本來就跨 run，逐 run 改會漏
            for run in list(para.runs)[1:]:
                run._element.getparent().remove(run._element)
            if para.runs:
                para.runs[0].text = text
            elif text:
                para.add_run(text)

    def walk(container):
        clean(container.paragraphs)
        for table in container.tables:
            for row in table.rows:
                for cell in row.cells:
                    walk(cell)          # 儲存格裡還可能有巢狀表格

    walk(doc)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _swap_images(value, tpl):
    """把 context 裡的圖片佔位標記換成 docxtpl 的 InlineImage。

    對照表無法自己建 InlineImage（那需要 tpl 物件），所以改成放
    `{'__image__': <bytes>, 'max_width_cm': 14, 'max_height_cm': 10}` 這種標記。

    尺寸照 EAGLE 原系統 imageGenerator 的算法**等比縮放**（上限 14×10 cm），
    先前寫死高度 60mm 會把直式照片壓扁。
    """
    from docx.shared import Cm
    from docxtpl import InlineImage

    from . import photo_stamp

    if isinstance(value, dict):
        if IMAGE_KEY in value:
            raw = value.get(IMAGE_KEY)
            if not raw:
                return ''
            width_px, height_px = photo_stamp.image_size_px(raw)
            width_cm, height_cm = photo_stamp.fit_size_cm(
                width_px, height_px,
                value.get('max_width_cm', photo_stamp.DEFAULT_MAX_WIDTH_CM),
                value.get('max_height_cm', photo_stamp.DEFAULT_MAX_HEIGHT_CM))
            return InlineImage(tpl, io.BytesIO(raw),
                               width=Cm(width_cm), height=Cm(height_cm))
        return {k: _swap_images(v, tpl) for k, v in value.items()}
    if isinstance(value, list):
        return [_swap_images(v, tpl) for v in value]
    return value


def render(raw_bytes, context):
    """把 context 套進 docx 樣板，回傳新的 docx bytes。"""
    from docxtpl import DocxTemplate
    # 匯入既有轉換器（dobtor_doc_editor 已在 depends 裡）
    from odoo.addons.dobtor_doc_editor.controllers.doc_controller import (
        _convert_ins_to_jinja)

    stripped = _strip_conditionals(raw_bytes)
    converted = _convert_ins_to_jinja(stripped)

    tpl = DocxTemplate(io.BytesIO(converted))
    tpl.render(_swap_images(context, tpl))
    out = io.BytesIO()
    tpl.save(out)
    return out.getvalue()


def _chunks(records, page_size):
    if not records:
        return [[]]
    return [records[i:i + page_size] for i in range(0, len(records), page_size)]


def paginate_plain(records, page_size):
    """`+++FOR item IN $page+++` 用——page 本身就是清單。

    送審管制、檢試驗管制的樣板是這種寫法。
    """
    return _chunks(records, page_size)


def paginate_pages(records, page_size, key='errorRecords'):
    """`+++FOR record IN $page.errorRecords+++` 用——page 是物件。

    缺失改善管制表的樣板是這種寫法，且要 $page.currentPage 印頁碼。
    """
    return [{'currentPage': i + 1, key: chunk}
            for i, chunk in enumerate(_chunks(records, page_size))]
